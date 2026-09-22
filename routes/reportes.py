from flask import Blueprint, request, jsonify, send_file
from flask_jwt_extended import jwt_required
from utils.decorators import admin_required
from models import db, Pedido, DetallePedido, Producto, Usuario, InventarioMovimiento, CatalogoEstadoPedido, CatalogoRolUsuario
from sqlalchemy import func, extract
from datetime import datetime
import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

reportes_bp = Blueprint('reportes', __name__, url_prefix='/api/reportes')

def _excel_response(wb, filename):
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(output, as_attachment=True, download_name=filename, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

def _generar_excel(titulo, headers, rows, filename, subtitulo=None):
    wb = Workbook()
    ws = wb.active
    ws.title = titulo[:30]
    fill = PatternFill(start_color="8CB07A", end_color="8CB07A", fill_type="solid")
    font = Font(color="FFFFFF", bold=True)
    ws.append([titulo])
    # Merge solo si hay headers
    if headers:
        end_col = chr(64 + len(headers))
        ws.merge_cells(f'A1:{end_col}1')
        ws['A1'].font = Font(size=14, bold=True, color="3A3E40")
        ws['A1'].alignment = Alignment(horizontal='center')
        # Subtítulo con periodo si se proporciona
        header_row = 2
        if subtitulo:
            ws.append([subtitulo])
            ws.merge_cells(f'A2:{end_col}2')
            ws['A2'].font = Font(size=10, italic=True, color="565659")
            ws['A2'].alignment = Alignment(horizontal='center')
            header_row = 3
            ws.append(headers)
        else:
            ws.append(headers)
        # Estilo encabezado
        for col in range(1, len(headers)+1):
            cell = ws.cell(row=header_row, column=col)
            cell.fill = fill
            cell.font = font
            cell.alignment = Alignment(horizontal='center')
        for row in rows:
            # Asegurar que row tenga misma longitud que headers (rellenar vacíos)
            if len(row) < len(headers):
                row = list(row) + [''] * (len(headers) - len(row))
            ws.append(row)
        # Ajustar ancho sin usar ws.columns (evita MergedCell)
        for idx, h in enumerate(headers, start=1):
            max_len = len(str(h))
            for r in rows:
                if idx-1 < len(r) and r[idx-1] not in (None, ''):
                    max_len = max(max_len, len(str(r[idx-1])))
            col_letter = chr(64 + idx)
            ws.column_dimensions[col_letter].width = min(max_len + 4, 40)
    else:
        for row in rows:
            ws.append(row)
    return _excel_response(wb, filename)

def _parse_fecha(fecha_str):
    if not fecha_str:
        return None
    try:
        return datetime.strptime(fecha_str, '%Y-%m-%d')
    except:
        return None

@reportes_bp.route('/ventas-totales', methods=['GET'])
@jwt_required()
@admin_required
def ventas_totales():
    try:
        desde = _parse_fecha(request.args.get('desde'))
        hasta = _parse_fecha(request.args.get('hasta'))
        excel = request.args.get('excel') == '1' or request.args.get('formato') == 'excel'

        estado_cancelado = CatalogoEstadoPedido.por_nombre('cancelado')
        query = Pedido.query.filter(Pedido.estado_id != estado_cancelado.id)
        if desde:
            query = query.filter(Pedido.fecha_pedido >= desde)
        if hasta:
            # Incluir todo el día
            hasta_fin = hasta.replace(hour=23, minute=59, second=59)
            query = query.filter(Pedido.fecha_pedido <= hasta_fin)

        total_pedidos = query.count()
        total_ventas = query.with_entities(func.coalesce(func.sum(Pedido.monto_total), 0)).scalar() or 0

        if excel:
            headers = ['Métrica', 'Valor']
            rows = [
                ['Total Pedidos', total_pedidos],
                ['Total Ventas (Q)', float(total_ventas)],
            ]
            if desde or hasta:
                rows.append(['Desde', desde.strftime('%Y-%m-%d') if desde else ''])
                rows.append(['Hasta', hasta.strftime('%Y-%m-%d') if hasta else ''])
            return _generar_excel('Ventas Totales', headers, rows, 'ventas_totales.xlsx')

        return jsonify({
            'data': {
                'total_pedidos': total_pedidos,
                'total_ventas': float(total_ventas),
                'desde': desde.isoformat() if desde else None,
                'hasta': hasta.isoformat() if hasta else None
            },
            'message': 'Ventas totales obtenidas'
        }), 200
    except Exception as e:
        return jsonify({'error': 'Error', 'message': str(e)}), 500

@reportes_bp.route('/ventas-por-mes', methods=['GET'])
@jwt_required()
@admin_required
def ventas_por_mes():
    try:
        excel = request.args.get('excel') == '1' or request.args.get('formato') == 'excel'
        # Agrupar por año-mes
        estado_cancelado = CatalogoEstadoPedido.por_nombre('cancelado')
        resultados = db.session.query(
            extract('year', Pedido.fecha_pedido).label('anio'),
            extract('month', Pedido.fecha_pedido).label('mes'),
            func.count(Pedido.id).label('pedidos'),
            func.sum(Pedido.monto_total).label('total')
        ).filter(Pedido.estado_id != estado_cancelado.id).group_by('anio', 'mes').order_by('anio', 'mes').all()

        data = [{'anio': int(r.anio), 'mes': int(r.mes), 'pedidos': r.pedidos, 'total': float(r.total or 0)} for r in resultados]

        if excel:
            headers = ['Año', 'Mes', 'Pedidos', 'Total Ventas (Q)']
            rows = [[d['anio'], d['mes'], d['pedidos'], d['total']] for d in data]
            return _generar_excel('Ventas por Mes', headers, rows, 'ventas_por_mes.xlsx')

        return jsonify({'data': data, 'message': 'Ventas por mes obtenidas'}), 200
    except Exception as e:
        return jsonify({'error': 'Error', 'message': str(e)}), 500

@reportes_bp.route('/ventas-por-periodo', methods=['GET'])
@jwt_required()
@admin_required
def ventas_por_periodo():
    try:
        desde_str = request.args.get('desde')
        hasta_str = request.args.get('hasta')
        excel = request.args.get('excel') == '1' or request.args.get('formato') == 'excel'
        if not desde_str or not hasta_str:
            return jsonify({'error': 'Fechas requeridas', 'message': 'Se requieren desde y hasta (YYYY-MM-DD)'}), 400
        desde = _parse_fecha(desde_str)
        hasta = _parse_fecha(hasta_str)
        if not desde or not hasta:
            return jsonify({'error': 'Formato inválido', 'message': 'Use YYYY-MM-DD'}), 400
        hasta_fin = hasta.replace(hour=23, minute=59, second=59)
        estado_cancelado = CatalogoEstadoPedido.por_nombre('cancelado')
        pedidos = Pedido.query.filter(Pedido.fecha_pedido >= desde, Pedido.fecha_pedido <= hasta_fin, Pedido.estado_id != estado_cancelado.id).order_by(Pedido.fecha_pedido.desc()).all()
        total = sum(float(p.monto_total) for p in pedidos)

        if excel:
            headers = ['ID', 'Fecha', 'Cliente', 'Estado', 'Total (Q)']
            rows = []
            for p in pedidos:
                nombre = p.usuario.nombre if p.usuario else (p.email_contacto or 'Invitado')
                rows.append([p.id, p.fecha_pedido.strftime('%Y-%m-%d'), nombre, p.estado_nombre, float(p.monto_total)])
            rows.append([])
            rows.append(['Total', '', '', '', total])
            subtitulo = f'Periodo: {desde_str} al {hasta_str}'
            return _generar_excel('Ventas por Periodo', headers, rows, f'ventas_{desde_str}_{hasta_str}.xlsx', subtitulo=subtitulo)

        return jsonify({
            'data': {
                'pedidos': [p.to_dict() for p in pedidos],
                'total_ventas': total,
                'cantidad': len(pedidos),
                'desde': desde_str,
                'hasta': hasta_str
            },
            'message': 'Ventas por periodo obtenidas'
        }), 200
    except Exception as e:
        return jsonify({'error': 'Error', 'message': str(e)}), 500

@reportes_bp.route('/ganancias', methods=['GET'])
@jwt_required()
@admin_required
def ganancias():
    try:
        desde = _parse_fecha(request.args.get('desde'))
        hasta = _parse_fecha(request.args.get('hasta'))
        excel = request.args.get('excel') == '1' or request.args.get('formato') == 'excel'

        # Ventas (ingresos)
        estado_cancelado = CatalogoEstadoPedido.por_nombre('cancelado')
        q_ventas = Pedido.query.filter(Pedido.estado_id != estado_cancelado.id)
        if desde:
            q_ventas = q_ventas.filter(Pedido.fecha_pedido >= desde)
        if hasta:
            hasta_fin = hasta.replace(hour=23, minute=59, second=59)
            q_ventas = q_ventas.filter(Pedido.fecha_pedido <= hasta_fin)
        ingresos = q_ventas.with_entities(func.coalesce(func.sum(Pedido.monto_total), 0)).scalar() or 0
        ingresos = float(ingresos)

        # Costos (entradas con costo_unitario)
        q_costos = InventarioMovimiento.query.filter_by(tipo='entrada').filter(InventarioMovimiento.costo_unitario.isnot(None))
        if desde:
            q_costos = q_costos.filter(InventarioMovimiento.fecha >= desde)
        if hasta:
            q_costos = q_costos.filter(InventarioMovimiento.fecha <= hasta_fin)
        costos_data = q_costos.all()
        costos = sum(float(m.costo_unitario * m.cantidad) for m in costos_data) if costos_data else 0

        ganancia = ingresos - costos
        margen = (ganancia / ingresos * 100) if ingresos > 0 else 0

        if excel:
            headers = ['Concepto', 'Monto (Q)']
            rows = [
                ['Ingresos (Ventas)', ingresos],
                ['Costos (Entradas)', costos],
                ['Ganancia', ganancia],
                ['Margen %', round(margen, 2)],
            ]
            return _generar_excel('Ganancias', headers, rows, 'ganancias.xlsx')

        return jsonify({
            'data': {
                'ingresos': ingresos,
                'costos': costos,
                'ganancia': ganancia,
                'margen_porcentaje': round(margen, 2),
                'desde': desde.isoformat() if desde else None,
                'hasta': hasta.isoformat() if hasta else None
            },
            'message': 'Ganancias obtenidas'
        }), 200
    except Exception as e:
        return jsonify({'error': 'Error', 'message': str(e)}), 500

@reportes_bp.route('/productos-mas-vendidos', methods=['GET'])
@jwt_required()
@admin_required
def productos_mas_vendidos():
    try:
        limit = request.args.get('limit', 10, type=int)
        desde = _parse_fecha(request.args.get('desde'))
        hasta = _parse_fecha(request.args.get('hasta'))
        excel = request.args.get('excel') == '1' or request.args.get('formato') == 'excel'

        estado_cancelado = CatalogoEstadoPedido.por_nombre('cancelado')
        q = db.session.query(
            Producto.id,
            Producto.nombre,
            Producto.marca_id,
            func.sum(DetallePedido.cantidad).label('total_vendidos'),
            func.sum(DetallePedido.subtotal).label('total_ingresos')
        ).join(DetallePedido, DetallePedido.producto_id == Producto.id)\
         .join(Pedido, Pedido.id == DetallePedido.pedido_id)\
         .filter(Pedido.estado_id != estado_cancelado.id)

        if desde:
            q = q.filter(Pedido.fecha_pedido >= desde)
        if hasta:
            hasta_fin = hasta.replace(hour=23, minute=59, second=59)
            q = q.filter(Pedido.fecha_pedido <= hasta_fin)

        q = q.group_by(Producto.id, Producto.nombre, Producto.marca_id).order_by(func.sum(DetallePedido.cantidad).desc()).limit(limit)
        resultados = q.all()

        data = []
        for r in resultados:
            marca = db.session.get(Producto, r.id).marca
            data.append({
                'producto_id': r.id,
                'nombre': r.nombre,
                'marca': marca.nombre if marca else '',
                'total_vendidos': int(r.total_vendidos),
                'total_ingresos': float(r.total_ingresos or 0)
            })

        if excel:
            headers = ['Producto', 'Marca', 'Unidades Vendidas', 'Ingresos (Q)']
            rows = [[d['nombre'], d['marca'], d['total_vendidos'], d['total_ingresos']] for d in data]
            return _generar_excel('Productos Mas Vendidos', headers, rows, 'productos_mas_vendidos.xlsx')

        return jsonify({'data': data, 'message': 'Productos más vendidos obtenidos'}), 200
    except Exception as e:
        return jsonify({'error': 'Error', 'message': str(e)}), 500

@reportes_bp.route('/clientes-top', methods=['GET'])
@jwt_required()
@admin_required
def clientes_top():
    try:
        limit = request.args.get('limit', 10, type=int)
        desde = _parse_fecha(request.args.get('desde'))
        hasta = _parse_fecha(request.args.get('hasta'))
        excel = request.args.get('excel') == '1' or request.args.get('formato') == 'excel'

        rol_cliente = CatalogoRolUsuario.por_nombre('cliente')
        estado_cancelado = CatalogoEstadoPedido.por_nombre('cancelado')
        q = db.session.query(
            Usuario.id,
            Usuario.nombre,
            Usuario.email,
            func.count(Pedido.id).label('total_pedidos'),
            func.sum(Pedido.monto_total).label('total_gastado')
        ).join(Pedido, Pedido.usuario_id == Usuario.id)\
         .filter(Usuario.rol_id == rol_cliente.id).filter(Pedido.estado_id != estado_cancelado.id)

        if desde:
            q = q.filter(Pedido.fecha_pedido >= desde)
        if hasta:
            hasta_fin = hasta.replace(hour=23, minute=59, second=59)
            q = q.filter(Pedido.fecha_pedido <= hasta_fin)

        q = q.group_by(Usuario.id, Usuario.nombre, Usuario.email).order_by(func.sum(Pedido.monto_total).desc()).limit(limit)
        resultados = q.all()

        data = []
        for r in resultados:
            data.append({
                'usuario_id': r.id,
                'nombre': r.nombre,
                'email': r.email,
                'total_pedidos': int(r.total_pedidos),
                'total_gastado': float(r.total_gastado or 0)
            })

        if excel:
            headers = ['Cliente', 'Email', 'Pedidos', 'Total Gastado (Q)']
            rows = [[d['nombre'], d['email'], d['total_pedidos'], d['total_gastado']] for d in data]
            return _generar_excel('Clientes Top', headers, rows, 'clientes_top.xlsx')

        return jsonify({'data': data, 'message': 'Clientes top obtenidos'}), 200
    except Exception as e:
        return jsonify({'error': 'Error', 'message': str(e)}), 500
